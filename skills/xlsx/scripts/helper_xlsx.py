#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Helper script for Excel (.xlsx) manipulation in Antigravity 2.0.
"""

import sys
import subprocess

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

def read_xlsx(file_path, sheet_name=None, data_only=True):
    """Read data from an Excel spreadsheet."""
    wb = openpyxl.load_workbook(file_path, data_only=data_only)
    sheet = wb[sheet_name] if sheet_name else wb.active
    rows = []
    for row in sheet.iter_rows(values_only=True):
        if any(cell is not None for cell in row):
            rows.append(list(row))
    return rows

def create_table_sheet(headers, data, output_path="table.xlsx", sheet_name="Data"):
    """Create a formatted table sheet."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    ws.append(headers)
    hdr_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    hdr_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    for col_idx in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col_idx)
        c.fill = hdr_fill
        c.font = hdr_font
        c.alignment = Alignment(horizontal="center")

    for row in data:
        ws.append(row)

    for col in ws.columns:
        max_len = max(len(str(c.value or '')) for c in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(output_path)
    return output_path

if __name__ == "__main__":
    if len(sys.argv) > 1:
        data = read_xlsx(sys.argv[1])
        print(f"Read {len(data)} rows from {sys.argv[1]}")
